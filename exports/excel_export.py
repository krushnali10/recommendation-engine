import os
import sys
import json
import sqlite3
import pandas as pd
import numpy as np
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

sys.path.insert(0, os.path.dirname(os.path.dirname(__file__)))
from models.engine import RecommendationEngine

# ── Colours ───────────────────────────────────────────────────────────────────
NAVY      = "1F3864"
BLUE      = "2E75B6"
WHITE     = "FFFFFF"
ALT_ROW   = "EBF3FB"
BORDER_C  = "B8CCE4"
GREEN_H   = "C6EFCE"
AMBER_H   = "FFEB9C"
RED_H     = "FFC7CE"


def _side(color=BORDER_C):
    s = Side(style="thin", color=color)
    return Border(left=s, right=s, top=s, bottom=s)


def _header(cell, bg=NAVY, fg=WHITE, size=11):
    cell.font      = Font(bold=True, color=fg, size=size, name="Arial")
    cell.fill      = PatternFill("solid", fgColor=bg)
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cell.border    = _side(bg)


def _data(cell, alt=False):
    cell.font      = Font(size=10, name="Arial")
    cell.fill      = PatternFill("solid", fgColor=ALT_ROW if alt else WHITE)
    cell.alignment = Alignment(vertical="center")
    cell.border    = _side()


def _widths(ws, mapping):
    for col, w in mapping.items():
        ws.column_dimensions[col].width = w


def _score_fill(score):
    if score >= 0.7: return GREEN_H
    if score >= 0.4: return AMBER_H
    return RED_H


# ── Sheet 1: Executive Summary ────────────────────────────────────────────────
def sheet_summary(wb, engine):
    ws = wb.active
    ws.title = "Executive Summary"
    ws.sheet_view.showGridLines = False

    # Title
    ws.merge_cells("A1:H1")
    c = ws["A1"]
    c.value     = "ML-Powered Recommendation Engine  —  Business Summary"
    c.font      = Font(bold=True, size=15, color=WHITE, name="Arial")
    c.fill      = PatternFill("solid", fgColor=NAVY)
    c.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 34

    ws.merge_cells("A2:H2")
    s = ws["A2"]
    s.value     = (f"Generated: {datetime.now().strftime('%B %d, %Y  %H:%M')}  |  "
                   f"Dataset: Amazon Electronics Reviews  |  "
                   f"Model: Hybrid CF + TF-IDF")
    s.font      = Font(italic=True, size=10, color=WHITE, name="Arial")
    s.fill      = PatternFill("solid", fgColor=BLUE)
    s.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[2].height = 20

    # KPI tiles
    kpis = [
        ("Total Products",    f"{len(engine.products):,}",      "A", "B"),
        ("Total Users",       f"{len(engine.users):,}",         "C", "D"),
        ("Interactions",      f"{len(engine.interactions):,}",  "E", "F"),
        ("Avg Rating",        f"{engine.products['rating'].mean():.2f} / 5.0", "G", "H"),
    ]
    for label, value, c1, c2 in kpis:
        ws.merge_cells(f"{c1}4:{c2}4")
        ws.merge_cells(f"{c1}5:{c2}5")
        ws.merge_cells(f"{c1}6:{c2}6")
        lbl = ws[f"{c1}4"]
        lbl.value     = label
        lbl.font      = Font(bold=True, size=9, color=WHITE, name="Arial")
        lbl.fill      = PatternFill("solid", fgColor=BLUE)
        lbl.alignment = Alignment(horizontal="center")
        val = ws[f"{c1}5"]
        val.value     = value
        val.font      = Font(bold=True, size=20, color=NAVY, name="Arial")
        val.fill      = PatternFill("solid", fgColor=ALT_ROW)
        val.alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[5].height = 36

    # Category table
    row = 9
    ws.merge_cells(f"A{row}:D{row}")
    h = ws[f"A{row}"]
    h.value = "Products by Category"
    _header(h, bg=BLUE)
    ws.row_dimensions[row].height = 20

    for ci, hdr in enumerate(["Category", "Count", "Share %", "Avg Rating"], 1):
        _header(ws.cell(row+1, ci, hdr), bg=BLUE)

    cat_stats = (engine.products.groupby("category")
                 .agg(count=("product_id","count"), avg_rating=("rating","mean"))
                 .reset_index())
    cat_stats["share"] = (cat_stats["count"] / len(engine.products) * 100).round(1)

    for i, r in enumerate(cat_stats.itertuples(), row+2):
        alt = i % 2 == 0
        for ci, v in enumerate([r.category, r.count, f"{r.share}%",
                                 f"{r.avg_rating:.2f}"], 1):
            _data(ws.cell(i, ci, v), alt)

    # Feature engineering table
    row2 = 9
    ws.merge_cells(f"F{row2}:H{row2}")
    h2 = ws[f"F{row2}"]
    h2.value = "Hybrid Score Weights"
    _header(h2, bg=BLUE)

    for ci, hdr in enumerate(["Signal", "Weight", "Description"], 6):
        _header(ws.cell(row2+1, ci, hdr), bg=BLUE)

    fe = [
        ("User-based CF",   "35%", "Cosine similarity between users"),
        ("Item-based CF",   "30%", "Item-item interaction similarity"),
        ("TF-IDF Content",  "25%", "Tag & category text similarity"),
        ("Popularity",      " 7%", "Log-scaled engagement score"),
        ("Quality",         " 3%", "Rating x log(reviews)"),
        ("Recency Boost",   " — ", "1.5x (<30d), 1.2x (<90d)"),
        ("Event Weight",    " — ", "Purchase=5x, Cart=3x, View=1x"),
    ]
    for i, (sig, wt, desc) in enumerate(fe, row2+2):
        for ci, v in enumerate([sig, wt, desc], 6):
            _data(ws.cell(i, ci, v), i % 2 == 0)

    _widths(ws, {"A":22,"B":10,"C":10,"D":12,"E":4,"F":18,"G":10,"H":28})


# ── Sheet 2: Top-N Recommendations ───────────────────────────────────────────
def sheet_recommendations(wb, engine, user_ids, top_n=10):
    ws = wb.create_sheet("Top-N Recommendations")
    ws.sheet_view.showGridLines = False

    ws.merge_cells("A1:I1")
    c = ws["A1"]
    c.value     = f"Top-{top_n} Personalised Recommendations — Sample Users"
    c.font      = Font(bold=True, size=13, color=WHITE, name="Arial")
    c.fill      = PatternFill("solid", fgColor=NAVY)
    c.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 28

    headers = ["User ID","Rank","Product ID","Category","Brand",
               "Price ($)","Rating","Score","Method"]
    for ci, h in enumerate(headers, 1):
        _header(ws.cell(2, ci, h))
    ws.row_dimensions[2].height = 18
    ws.freeze_panes = "A3"

    row = 3
    for uid in user_ids:
        try:
            recs = engine.recommend(uid, top_n=top_n, method="hybrid")
        except Exception:
            continue
        for rank, (_, r) in enumerate(recs.iterrows(), 1):
            alt = row % 2 == 0
            vals = [uid, rank, r["product_id"], r["category"], r["brand"],
                    r["price"], r["rating"], r["recommendation_score"], r["method"]]
            for ci, v in enumerate(vals, 1):
                cell = ws.cell(row, ci, v)
                _data(cell, alt)
                if ci == 6: cell.number_format = "#,##0.00"
                if ci == 8:
                    cell.number_format = "0.0000"
                    cell.fill = PatternFill("solid", fgColor=_score_fill(float(v)))
            ws.row_dimensions[row].height = 15
            row += 1

    _widths(ws, {"A":20,"B":6,"C":14,"D":14,"E":12,"F":10,"G":8,"H":10,"I":14})


# ── Sheet 3: Similar Products ─────────────────────────────────────────────────
def sheet_similarity(wb, engine, product_ids, top_n=8):
    ws = wb.create_sheet("Product Similarity")
    ws.sheet_view.showGridLines = False

    ws.merge_cells("A1:H1")
    c = ws["A1"]
    c.value     = "Item-to-Item Similarity (Content-Based TF-IDF)"
    c.font      = Font(bold=True, size=13, color=WHITE, name="Arial")
    c.fill      = PatternFill("solid", fgColor=NAVY)
    c.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 28

    headers = ["Source Product","Source Category","Similar Product",
               "Category","Brand","Price ($)","Rating","Similarity Score"]
    for ci, h in enumerate(headers, 1):
        _header(ws.cell(2, ci, h))
    ws.freeze_panes = "A3"

    row = 3
    for pid in product_ids:
        src  = engine.products[engine.products["product_id"] == pid]
        if src.empty: continue
        src  = src.iloc[0]
        sims = engine.recommend_similar(pid, top_n=top_n)
        for _, r in sims.iterrows():
            alt = row % 2 == 0
            vals = [src["product_id"], src["category"], r["product_id"],
                    r["category"], r["brand"], r["price"], r["rating"],
                    r["similarity_score"]]
            for ci, v in enumerate(vals, 1):
                cell = ws.cell(row, ci, v)
                _data(cell, alt)
                if ci == 6: cell.number_format = "#,##0.00"
                if ci == 8:
                    cell.number_format = "0.0000"
                    cell.fill = PatternFill("solid", fgColor=_score_fill(float(v)))
            ws.row_dimensions[row].height = 15
            row += 1

    _widths(ws, {"A":16,"B":14,"C":16,"D":14,"E":12,"F":10,"G":8,"H":16})


# ── Sheet 4: Interaction Analytics ───────────────────────────────────────────
def sheet_analytics(wb, engine):
    ws = wb.create_sheet("Interaction Analytics")
    ws.sheet_view.showGridLines = False

    ws.merge_cells("A1:F1")
    c = ws["A1"]
    c.value     = "Interaction Analytics & Dataset Overview"
    c.font      = Font(bold=True, size=13, color=WHITE, name="Arial")
    c.fill      = PatternFill("solid", fgColor=NAVY)
    c.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 28

    # Event distribution
    ws.merge_cells("A3:C3")
    _header(ws["A3"], bg=BLUE)
    ws["A3"].value = "Event Type Distribution"
    for ci, h in enumerate(["Event","Count","Share %"], 1):
        _header(ws.cell(4, ci, h), bg=BLUE)

    ev = engine.interactions["event_type"].value_counts()
    for i, (evt, cnt) in enumerate(ev.items(), 5):
        for ci, v in enumerate([evt, cnt, f"{cnt/len(engine.interactions)*100:.1f}%"], 1):
            _data(ws.cell(i, ci, v), i % 2 == 0)

    # Rating distribution
    ws.merge_cells("E3:F3")
    _header(ws["E3"], bg=BLUE)
    ws["E3"].value = "Rating Distribution"
    for ci, h in enumerate(["Rating","Count"], 5):
        _header(ws.cell(4, ci, h), bg=BLUE)

    rd = engine.interactions["rating"].value_counts().sort_index()
    for i, (rating, cnt) in enumerate(rd.items(), 5):
        for ci, v in enumerate([rating, cnt], 5):
            _data(ws.cell(i, ci, v), i % 2 == 0)

    # Top 10 most interacted products
    row = 14
    ws.merge_cells(f"A{row}:F{row}")
    _header(ws[f"A{row}"], bg=BLUE)
    ws[f"A{row}"].value = "Top 10 Most Interacted Products"
    for ci, h in enumerate(["Product ID","Category","Brand","Price","Rating","Interactions"], 1):
        _header(ws.cell(row+1, ci, h), bg=BLUE)

    top10 = (engine.interactions.groupby("product_id").size()
             .nlargest(10).reset_index(name="count"))
    top10 = top10.merge(engine.products[["product_id","category","brand","price","rating"]],
                        on="product_id")
    for i, r in enumerate(top10.itertuples(), row+2):
        alt = i % 2 == 0
        for ci, v in enumerate([r.product_id, r.category, r.brand,
                                 r.price, r.rating, r.count], 1):
            cell = ws.cell(i, ci, v)
            _data(cell, alt)
            if ci == 4: cell.number_format = "#,##0.00"

    _widths(ws, {"A":16,"B":14,"C":12,"D":10,"E":8,"F":14})


# ── Sheet 5: Category Heatmap ─────────────────────────────────────────────────
def sheet_heatmap(wb, engine):
    ws = wb.create_sheet("Category Heatmap")
    ws.sheet_view.showGridLines = False

    ws.merge_cells("A1:I1")
    c = ws["A1"]
    c.value     = "Category x Event Type Interaction Heatmap"
    c.font      = Font(bold=True, size=13, color=WHITE, name="Arial")
    c.fill      = PatternFill("solid", fgColor=NAVY)
    c.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 28

    ix = engine.interactions.merge(
        engine.products[["product_id","category"]], on="product_id"
    )
    pivot = pd.crosstab(ix["category"], ix["event_type"])

    ws.cell(2, 1, "Category").font = Font(bold=True, name="Arial")
    for ci, col in enumerate(pivot.columns, 2):
        _header(ws.cell(2, ci, col), bg=BLUE)
    _header(ws.cell(2, len(pivot.columns)+2, "Total"), bg=NAVY)

    vmax = pivot.values.max()
    for ri, (cat, row_d) in enumerate(pivot.iterrows(), 3):
        ws.cell(ri, 1, cat).font = Font(bold=True, name="Arial")
        for ci, v in enumerate(row_d, 2):
            cell = ws.cell(ri, ci, int(v))
            intensity = int(v / vmax * 180)
            r_hex = format(255 - intensity//2, "02X")
            g_hex = format(min(255, 190 + intensity//6), "02X")
            b_hex = format(255 - intensity, "02X")
            cell.fill      = PatternFill("solid", fgColor=f"{r_hex}{g_hex}{b_hex}")
            cell.alignment = Alignment(horizontal="center")
            cell.border    = _side()
        total = ws.cell(ri, len(pivot.columns)+2, int(row_d.sum()))
        total.font      = Font(bold=True, name="Arial")
        total.fill      = PatternFill("solid", fgColor=ALT_ROW)
        total.alignment = Alignment(horizontal="center")

    for ci in range(1, len(pivot.columns)+3):
        ws.column_dimensions[get_column_letter(ci)].width = 16


# ── Main ──────────────────────────────────────────────────────────────────────
def export(engine, output_path, sample_users=20, sample_products=15, top_n=10):
    print("[Export] Building Excel report...")
    wb = Workbook()

    user_ids    = engine.users["user_id"].sample(
        min(sample_users, len(engine.users)), random_state=42).tolist()
    product_ids = engine.products["product_id"].sample(
        min(sample_products, len(engine.products)), random_state=42).tolist()

    print("  Sheet 1: Executive Summary")
    sheet_summary(wb, engine)

    print("  Sheet 2: Top-N Recommendations")
    sheet_recommendations(wb, engine, user_ids, top_n)

    print("  Sheet 3: Product Similarity")
    sheet_similarity(wb, engine, product_ids)

    print("  Sheet 4: Interaction Analytics")
    sheet_analytics(wb, engine)

    print("  Sheet 5: Category Heatmap")
    sheet_heatmap(wb, engine)

    wb.save(output_path)
    print(f"\n[Export] Saved to: {output_path}")


if __name__ == "__main__":
    BASE = os.path.dirname(os.path.dirname(__file__))
    DB   = os.path.join(BASE, "data", "recommendation.db")
    OUT  = os.path.join(BASE, "exports", "recommendation_report.xlsx")
    os.makedirs(os.path.dirname(OUT), exist_ok=True)

    eng = RecommendationEngine(DB)
    export(eng, OUT, sample_users=20, sample_products=15, top_n=10)